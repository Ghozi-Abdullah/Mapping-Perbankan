"""Pembaca tabel SEKI (Statistik Ekonomi dan Keuangan Indonesia) Bank Indonesia.

Tabel SEKI berbentuk lembar Excel lama (.xls) dengan baris judul, baris tahun, baris bulan, lalu baris data
berlabel di kiri (Indonesia) dan kanan (Inggris). Pembaca ini tidak bergantung pada nomor baris atau kolom:
ia mencari baris bulan, menentukan tahun tiap kolom, lalu membaca label dan angka tiap baris.

Satu berkas bisa berisi beberapa lembar: riwayat lama (misalnya klasifikasi lapangan usaha sebelum 2017) dan
lembar terbaru. Semua lembar dibaca, lalu digabung dari yang terlama ke yang terbaru supaya bulan terbaru selalu
masuk dan angka revisi di lembar terbaru yang dipakai.
"""

import html.parser
import io
import re
from datetime import date, datetime

from fetchlib import to_float

MONTHS = {
    "jan": 1, "januari": 1, "january": 1, "feb": 2, "februari": 2, "february": 2, "peb": 2, "mar": 3, "maret": 3, "march": 3,
    "apr": 4, "april": 4, "mei": 5, "may": 5, "jun": 6, "juni": 6, "june": 6, "jul": 7, "juli": 7, "july": 7,
    "agu": 8, "agt": 8, "ags": 8, "aug": 8, "agus": 8, "agustus": 8, "august": 8, "sep": 9, "sept": 9, "september": 9,
    "okt": 10, "oct": 10, "oktober": 10, "october": 10, "nov": 11, "nop": 11, "november": 11, "nopember": 11,
    "des": 12, "dec": 12, "desember": 12, "december": 12,
}
MONTH_RE = re.compile(r"^([a-z]{3,9})\b\.?\s*[\*\^]*\s*(?:[a-z]\))?\s*[\*\^]*\s*[-'/ ]?\s*(\d{2}|\d{4})?\s*[\*\^]*\s*(?:[a-z]\))?$")
YEAR_RE = re.compile(r"^\s*((?:19|20)\d{2})\s*[\*\^]*\s*(?:[a-z]\))?\s*$")

GROUPS = [
    ("persero", re.compile(r"bank persero|state(-owned)? banks?|bank bumn", re.I)),
    ("bpd", re.compile(r"pemerintah daerah|pembangunan daerah|regional (government|development)", re.I)),
    ("busn", re.compile(r"swasta nasional|private national|private banks?", re.I)),
    ("kcba", re.compile(r"luar negeri|foreign banks?|bank asing|campuran|joint venture", re.I)),
    ("bpr", re.compile(r"perkreditan rakyat|perekonomian rakyat|rural banks?", re.I)),
    ("umum", re.compile(r"^bank umum$|^commercial banks?$", re.I)),
    ("total", re.compile(r"^(jumlah|total)\b", re.I)),
]


def month_token(cell):
    """(tahun atau None, bulan) dari isi sel header; None bila bukan bulan."""
    if isinstance(cell, (date, datetime)):
        return cell.year, cell.month
    s = str(cell or "").strip().lower()
    if not s:
        return None
    m = MONTH_RE.match(s)
    if not m or m.group(1) not in MONTHS:
        return None
    year = m.group(2)
    if year:
        year = int(year) + (2000 if len(year) == 2 else 0)
    return (year or None), MONTHS[m.group(1)]


def year_token(cell):
    if isinstance(cell, (int, float)) and not isinstance(cell, bool) and 1990 <= cell <= 2100 and float(cell).is_integer():
        return int(cell)
    m = YEAR_RE.match(str(cell or ""))
    return int(m.group(1)) if m else None


def clean_label(s):
    s = re.sub(r"\s+", " ", str(s or "")).strip()
    s = re.sub(r"^(\d+[.)]\s*|\d+\s+(?!(bulan|bln|months?|tahun|years?|hari|days?)\b)|[a-z][.)]\s+|[ivx]+[.)]\s+)", "", s, flags=re.I)
    s = re.sub(r"\s*(\d\)|\*+\)?|[a-z]\))$", "", s)
    return s.strip(" .:")


# ---------------------------------------------------------------- membaca berkas


class _Table(html.parser.HTMLParser):
    def __init__(self):
        super().__init__()
        self.rows, self.row, self.cell, self.in_cell = [], None, [], False

    def handle_starttag(self, tag, attrs):
        if tag == "tr":
            self.row = []
        elif tag in ("td", "th") and self.row is not None:
            self.in_cell, self.cell = True, []
            span = dict(attrs).get("colspan")
            self.span = int(span) if span and span.isdigit() else 1

    def handle_endtag(self, tag):
        if tag in ("td", "th") and self.row is not None and self.in_cell:
            self.row.append(" ".join("".join(self.cell).split()))
            self.row.extend([""] * (self.span - 1))
            self.in_cell = False
        elif tag == "tr" and self.row is not None:
            self.rows.append(self.row)
            self.row = None

    def handle_data(self, data):
        if self.in_cell:
            self.cell.append(data)


def latest_month(grid, today=None):
    """Bulan terbaru di header lembar ("YYYY-MM"), atau None bila lembar bukan tabel."""
    h = find_header(grid)
    if h is None:
        return None
    cols = month_columns(grid, h, today)
    return max(cols.values()) if cols else None


def _sheet_score(grid, today=None):
    """Urutan pilihan lembar: bulan terbaru dulu, lalu jumlah kolom bulan terbanyak."""
    h = find_header(grid)
    if h is None:
        return ("", 0)
    return (latest_month(grid, today) or "", sum(1 for c in grid[h] if month_token(c)))


def read_grid(raw, today=None):
    """Lembar dengan bulan terbaru sebagai list baris."""
    return max(read_grids(raw), key=lambda g: _sheet_score(g, today))


def sheet_rows(raw, today=None):
    """[(bulan terbaru, baris data)] untuk setiap lembar bertabel, urut dari lembar terlama ke terbaru."""
    out = []
    for grid in read_grids(raw):
        latest = latest_month(grid, today)
        if latest is None:
            continue
        try:
            out.append((latest, data_rows(grid, today)))
        except ValueError:
            continue
    if not out:
        raise ValueError("baris bulan tidak ditemukan di lembar mana pun")
    return sorted(out, key=lambda x: x[0])


def read_grids(raw):
    """Semua lembar sebagai list baris. Mendukung .xls, .xlsx, dan tabel HTML berakhiran .xls."""
    head = raw[:8]
    grids = []
    if head.startswith(b"\xd0\xcf\x11\xe0"):
        import xlrd

        book = xlrd.open_workbook(file_contents=raw)
        for sheet in book.sheets():
            grid = []
            for r in range(sheet.nrows):
                row = []
                for c in range(sheet.ncols):
                    cell = sheet.cell(r, c)
                    if cell.ctype == xlrd.XL_CELL_DATE:
                        try:
                            row.append(datetime(*xlrd.xldate_as_tuple(cell.value, book.datemode)[:6]))
                        except Exception:
                            row.append(cell.value)
                    else:
                        row.append(cell.value)
                grid.append(row)
            grids.append(grid)
    elif head.startswith(b"PK"):
        import openpyxl

        book = openpyxl.load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
        grids = [[list(r) for r in ws.iter_rows(values_only=True)] for ws in book.worksheets]
    else:
        text = raw.decode("utf-8", "replace") if b"<" in raw[:2000] else ""
        if "<t" not in text.lower():
            raise ValueError("format berkas tidak dikenal")
        p = _Table()
        p.feed(text)
        grids = [p.rows]
    if not grids:
        raise ValueError("berkas tidak berisi lembar")
    return grids


# ---------------------------------------------------------------- struktur tabel


def find_header(grid, scan=40):
    """Indeks baris bulan: baris dengan sel bulan terbanyak di bagian atas lembar."""
    best, best_n = None, 3
    for i, row in enumerate(grid[:scan]):
        n = sum(1 for c in row if month_token(c))
        if n > best_n:
            best, best_n = i, n
    return best


def month_columns(grid, h, today=None):
    """{kolom: "YYYY-MM"} untuk baris header h. Tahun diambil dari sel bulan itu sendiri, baris tahun di atasnya,
    atau diurutkan mundur dari bulan terakhir bila lembar tidak mencantumkan tahun."""
    today = today or date.today()
    row = grid[h]
    cols = [(c, month_token(v)) for c, v in enumerate(row)]
    cols = [(c, t) for c, t in cols if t]
    years = {}
    for r in range(max(0, h - 4), h):
        for c, v in enumerate(grid[r]):
            y = year_token(v)
            if y:
                years[c] = y
    year_cols = sorted(years)

    def year_above(c):
        ys = [years[yc] for yc in year_cols if yc <= c]
        return ys[-1] if ys else None

    out, prev, pending = {}, None, []
    for c, (own, m) in cols:
        y = own or year_above(c)
        if prev:
            if y is None or y < prev[0]:
                y = prev[0]
            if own is None and y == prev[0] and m <= prev[1]:
                y += 1
        if y is None:
            pending.append((c, m))
            continue
        prev = (y, m)
        out[c] = f"{y:04d}-{m:02d}"
    if out and pending:
        # Kolom awal tanpa tahun: hitung mundur dari kolom pertama yang sudah bertahun.
        y, m = map(int, out[min(out)].split("-"))
        for c, mm in reversed(pending):
            if mm >= m:
                y -= 1
            m = mm
            out[c] = f"{y:04d}-{m:02d}"
    if not out:
        last_m = cols[-1][1][1]
        y = today.year if last_m <= today.month else today.year - 1
        m = last_m
        for c, (_, mm) in reversed(cols):
            if mm > m:
                y -= 1
            m = mm
            out[c] = f"{y:04d}-{m:02d}"
    # Tabel tidak mungkin berisi bulan setelah hari ini; bila terjadi, tahun salah tebak satu.
    latest = max(out.values())
    if latest > today.strftime("%Y-%m"):
        out = {c: f"{int(v[:4]) - 1:04d}{v[4:]}" for c, v in out.items()}
    seen, uniq = set(), {}
    for c in sorted(out):
        if out[c] not in seen:
            seen.add(out[c])
            uniq[c] = out[c]
    return uniq


def data_rows(grid, today=None):
    """[(label, {bulan: nilai})] untuk semua baris di bawah header, termasuk baris judul kelompok tanpa angka."""
    h = find_header(grid)
    if h is None:
        raise ValueError("baris bulan tidak ditemukan")
    cols = month_columns(grid, h, today)
    if not cols:
        raise ValueError("kolom bulan tidak terbaca")
    first, last = min(cols), max(cols)
    rows = []
    for row in grid[h + 1:]:
        label = ""
        for c in range(0, min(first, len(row))):
            v = row[c]
            if isinstance(v, str) and re.search(r"[A-Za-z]{2}", v):
                label = v
                break
        if not label:
            for c in range(last + 1, len(row)):
                v = row[c]
                if isinstance(v, str) and re.search(r"[A-Za-z]{2}", v):
                    label = v
                    break
        label = clean_label(label)
        if not label:
            continue
        vals = {}
        for c, ym in cols.items():
            if c < len(row):
                x = to_float(row[c], decimal=".")
                if x is not None:
                    vals[ym] = x
        rows.append((label, vals))
    return rows


def group_of(label):
    for key, pat in GROUPS:
        if pat.search(label):
            return key
    return None


# ---------------------------------------------------------------- ekstraksi per jenis tabel


def by_group_totals(rows):
    """Tabel kredit menurut kelompok bank: baris pertama tiap kelompok adalah totalnya."""
    out = {}
    for label, vals in rows:
        g = group_of(label)
        if g and g not in out and vals:
            out[g] = vals
    return out


def _sum(*series):
    series = [s for s in series if s]
    if not series:
        return {}
    months = set.intersection(*(set(s) for s in series))
    return {m: sum(s[m] for s in series) for m in sorted(months)}


SIMPANAN_TYPES = [("giro", re.compile(r"^(giro|demand deposits?)$", re.I)),
                  ("tabungan", re.compile(r"^(tabungan|savings?( deposits?)?)$", re.I)),
                  ("deposito", re.compile(r"^(deposito|simpanan berjangka|time deposits?)$", re.I))]
CURRENCY = [("idr", re.compile(r"^rupiah$", re.I)),
            ("fx", re.compile(r"^(valuta asing|valas|foreign currenc(y|ies))$", re.I)),
            ("all", re.compile(r"^(jumlah|total|rupiah dan valuta asing|rupiah (dan|&) valas|rupiah and foreign currenc(y|ies))$", re.I))]


def simpanan_by_group(rows):
    """Tabel simpanan menurut kelompok bank dan jenis: blok Rupiah, Valas, (Jumlah) masing-masing berisi kelompok dan jenis."""
    data = {}
    cur, grp = None, None
    for label, vals in rows:
        c = next((k for k, p in CURRENCY if p.search(label)), None)
        if c and (c != "all" or grp is None or cur in ("idr", "fx")):
            cur, grp = c, None
            if vals:
                data.setdefault(cur, {}).setdefault("_total", {})["total"] = vals
            continue
        g = group_of(label)
        if g and g != "total":
            grp = g
            if vals:
                data.setdefault(cur or "all", {}).setdefault(grp, {})["total"] = vals
            continue
        t = next((k for k, p in SIMPANAN_TYPES if p.search(label)), None)
        if t and grp and vals:
            data.setdefault(cur or "all", {}).setdefault(grp, {})[t] = vals
    out = {}
    groups = {g for block in data.values() for g in block if g != "_total"}
    for g in groups:
        out[g] = {}
        for t in ("total", "giro", "tabungan", "deposito"):
            if g in data.get("all", {}) and t in data["all"][g]:
                out[g][t] = data["all"][g][t]
            else:
                s = _sum(data.get("idr", {}).get(g, {}).get(t), data.get("fx", {}).get(g, {}).get(t))
                if s:
                    out[g][t] = s
        if "total" not in out[g]:
            out[g]["total"] = _sum(*(out[g].get(t) for t in ("giro", "tabungan", "deposito")))
    return out


LOAN_TYPES = [("modal_kerja", re.compile(r"modal kerja|working capital", re.I)),
              ("investasi", re.compile(r"investasi|investment", re.I)),
              ("konsumsi", re.compile(r"konsumsi|consum", re.I))]
TENORS = [(k, re.compile(rf"^{k}\s*(bulan|bln|months?)$", re.I)) for k in ("1", "3", "6", "12", "24")]


def rates_by_group(rows, types):
    """Tabel suku bunga: judul kelompok diikuti beberapa baris jenis. Baris kelompok yang berangka disimpan sebagai 'rata'."""
    out, grp = {}, None
    for label, vals in rows:
        g = group_of(label)
        if g:
            grp = "umum" if g == "total" else g
            if vals:
                out.setdefault(grp, {})["rata"] = vals
            continue
        t = next((k for k, p in types if p.search(label)), None)
        if t and grp and vals and t not in out.get(grp, {}):
            out.setdefault(grp, {})[t] = vals
    return out


def by_province(rows, kode_fn):
    out, total = {}, None
    for label, vals in rows:
        if not vals:
            continue
        k = kode_fn(label)
        if k and k not in out:
            out[k] = vals
        elif total is None and re.match(r"^(jumlah|total|indonesia)\b", label, re.I):
            total = vals
    return out, total
